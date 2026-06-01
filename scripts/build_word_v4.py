"""
将 system_architecture_v4.html 的内容生成为 Word 文档。
- 框架图/流程图：嵌入Puppeteer截取的PNG
- 文字/表格/代码：python-docx原生生成
"""

import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

SCREENSHOT_DIR = Path("/tmp/docx_screenshots")
OUTPUT_PATH = Path("outputs/system_architecture_v4.docx")
PROVENANCE_JSON = Path("outputs/stage2_provenance_v2.json")
FEATURES_JSON = Path("outputs/stage2_features.json")

# ── 样式辅助 ──

def set_cell_bg(cell, hex_color):
    """设置单元格背景色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_styled_paragraph(doc, text, style=None, bold=False, font_size=None,
                         color=None, alignment=None, space_after=None, space_before=None):
    """添加带样式的段落"""
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if bold:
        run.bold = True
    if font_size:
        run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    return p


def add_code_block(doc, code_text):
    """添加代码块（灰色背景，等宽字体）"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    # 用表格模拟代码块背景
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_bg(cell, "1E293B")
    cp = cell.paragraphs[0]
    run = cp.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(226, 232, 240)
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)


def add_image_or_fallback(doc, img_name, fallback_text, width=Inches(6.2)):
    """嵌入截图，如果不存在则显示占位文字"""
    img_path = SCREENSHOT_DIR / img_name
    if img_path.exists():
        doc.add_picture(str(img_path), width=width)
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_styled_paragraph(doc, f"[图片缺失: {img_name}] {fallback_text}",
                             color=(150, 150, 150), font_size=9)


def make_table(doc, headers, rows, col_widths=None):
    """创建带格式的表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_bg(cell, "F0FDFA")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(19, 78, 74)

    # 数据行
    for ri, row_data in enumerate(rows):
        for ci, val in enumerate(row_data):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    # 列宽
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    return table


# ── 主文档构建 ──

def build_document():
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    # ============================================================
    # 封面/标题
    # ============================================================
    doc.add_paragraph()
    doc.add_paragraph()
    title = add_styled_paragraph(doc, "KnowFeat: Knowledge-Traceable Feature Engineering",
                                 bold=True, font_size=22,
                                 color=(15, 23, 42),
                                 alignment=WD_ALIGN_PARAGRAPH.CENTER)
    subtitle = add_styled_paragraph(doc, "LLM Agent 自动特征工程系统 — 技术方案 v4",
                                    font_size=14, color=(100, 116, 139),
                                    alignment=WD_ALIGN_PARAGRAPH.CENTER,
                                    space_after=6)
    tags = add_styled_paragraph(doc, "ReAct Agent  ·  Knowledge-Traceable  ·  Three-Level Verification  ·  SHAP Feedback  ·  AML",
                                font_size=9, color=(100, 116, 139),
                                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                                space_after=24)

    # ============================================================
    # 目录
    # ============================================================
    doc.add_heading("目录", level=1)
    toc_items = [
        "1. 核心思想",
        "2. Agent 架构总览",
        "3. Agent 的领域知识注入",
        "4. Agent 的工具箱",
        "5. 验证特征有效性（三级验证）",
        "6. Agent 的搜索策略",
        "7. Agent 的记忆与学习",
        "8. 收敛检测与终止",
        "9. 评价指标与评估方法",
        "10. 与传统方法的对比",
        "附. 一次完整交互示例",
        "R. 实验结果：18个特征知识溯源卡片",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)

    doc.add_page_break()

    # ============================================================
    # 第1章 核心思想
    # ============================================================
    doc.add_heading("1. 核心思想", level=1)

    add_styled_paragraph(doc,
        "传统特征工程方法（OpenFE、AutoFeat）通过数学组合搜索生成特征，效果好但缺乏语义解释。"
        "LLM方法（CAAFE、FeatLLM）利用大模型的语义理解能力，但生成的特征无法追溯到具体的领域知识来源。",
        font_size=10.5)

    add_styled_paragraph(doc,
        "KnowFeat 的核心创新：每个生成的特征都携带完整的知识溯源（Provenance Card），"
        "可以追溯到具体的风险指标（IND_xx）、检测规则（RULE_xxx）、专家意见（EXP_xx）和裁判文书（RF_S2_xxx）。",
        font_size=10.5, space_after=8)

    add_image_or_fallback(doc, "sec1_core_idea.png", "核心思想图")

    doc.add_heading("核心设计理念", level=2)
    items = [
        ("知识可追溯 (Knowledge-Traceable)", "每个特征 → 具体知识条目（IND/RULE/EXP/RF）→ 设计推理链"),
        ("Agent自主搜索", "ReAct范式，LLM通过工具调用自主探索数据、生成特征、验证效果"),
        ("三级验证", "L1沙盒执行 → L2统计质量(噪声校准) → L3模型效果(delta-metric)"),
        ("SHAP反馈驱动", "SHAP重要性分析 → 动态调整各维度探索预算 → 薄弱维度优先"),
    ]
    for title_text, desc in items:
        p = doc.add_paragraph()
        run_t = p.add_run(f"• {title_text}：")
        run_t.bold = True
        run_t.font.size = Pt(10.5)
        run_d = p.add_run(desc)
        run_d.font.size = Pt(10.5)

    doc.add_page_break()

    # ============================================================
    # 第2章 Agent架构总览
    # ============================================================
    doc.add_heading("2. Agent 架构总览", level=1)

    doc.add_heading("2.1 ReAct 范式", level=2)
    add_styled_paragraph(doc,
        "系统采用 ReAct（Reasoning + Acting）范式：Agent 在每一步先思考（Reasoning），"
        "再行动（通过 Tool Use 调用工具），最后根据观测（工具返回结果）决定下一步。"
        "LLM 通过 Function Calling 协议与环境交互。",
        font_size=10.5, space_after=8)

    # ReAct流程图截图
    add_image_or_fallback(doc, "flow_0.png", "ReAct单步循环流程图", width=Inches(5.5))

    p = doc.add_paragraph()
    run = p.add_run("[1] Yao, S. et al. \"ReAct: Synergizing Reasoning and Acting in Language Models.\" ICLR 2023.")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("2.2 完整架构图", level=2)
    add_styled_paragraph(doc,
        "系统分两个阶段运行：阶段一使用加密货币通用知识 + BABD-13数据集；"
        "阶段二注入数字人民币专属知识（38风险指标、15检测规则、26段专家意见、309份裁判文书）。"
        "两个阶段共享相同的Agent架构和工具箱。",
        font_size=10.5, space_after=8)

    add_image_or_fallback(doc, "arch_diagram.png", "完整架构图（两阶段）", width=Inches(6.2))

    add_styled_paragraph(doc,
        "图注：绿色节点为Agent（决策者），橙色节点为工具（执行者），紫色节点为SHAP反馈模块。"
        "阶段一产出Feature Set A作为基线列带入阶段二。最终输出有效特征集、Provenance Card和搜索日志。",
        font_size=9, color=(100, 116, 139), space_after=12)

    doc.add_page_break()

    # ============================================================
    # 第3章 领域知识注入
    # ============================================================
    doc.add_heading("3. Agent 的领域知识注入", level=1)

    add_image_or_fallback(doc, "sec3_knowledge_injection.png", "领域知识注入详情", width=Inches(6.2))

    doc.add_heading("3.1 知识资产体系", level=2)

    make_table(doc,
        ["知识类型", "ID格式", "数量", "来源"],
        [
            ["风险指标", "IND_01 ~ IND_38", "38项", "风险分类表_38指标.xlsx"],
            ["检测规则", "RULE_001 ~ RULE_015", "15条", "检测规则_15条.xlsx"],
            ["专家意见", "EXP_01 ~ EXP_26", "26段", "邮储专家访谈_数币洗钱模式.docx"],
            ["交易模式", "PAT_N01~N03, PAT_A01~A05", "8种", "系统硬编码（3正常+5异常）"],
            ["裁判文书风险因子", "RF_S2_FFF_NNN", "~1920条", "309份裁判文书JSON"],
        ],
        col_widths=[3, 4.5, 2, 5])

    doc.add_heading("3.2 Key-clue Extractor", level=2)
    add_styled_paragraph(doc,
        "Key-clue Extractor 将异构、非结构化的专家知识转换为统一的结构化schema：\n"
        "pattern（风险模式）/ dimension（维度标签）/ observable_signals（可观测信号）/ data_hint（数据字段提示）。\n"
        "309份裁判文书经过按维度聚合去重后压缩至几千字，注入LLM的System Prompt。",
        font_size=10.5, space_after=8)

    doc.add_heading("3.3 知识注入方式", level=2)
    add_styled_paragraph(doc,
        "所有领域知识在 System Prompt 中以 Markdown 格式注入，包含：\n"
        "① 数据Schema（字段名、类型、示例值）\n"
        "② 分布差异统计（异常vs正常各字段的均值/中位数/差异倍数）\n"
        "③ 带ID编号的知识条目（IND/RULE/EXP/PAT/RF）\n"
        "④ 溯源要求：必须在 knowledge_ids 中填写具体的知识资产ID",
        font_size=10.5)

    doc.add_page_break()

    # ============================================================
    # 第4章 工具箱
    # ============================================================
    doc.add_heading("4. Agent 的工具箱", level=1)

    add_image_or_fallback(doc, "sec4_toolbox.png", "Agent工具箱详情", width=Inches(6.2))

    make_table(doc,
        ["工具", "功能", "输入", "输出"],
        [
            ["explore_data", "数据查询与探索", "Python查询代码", "DataFrame结果（限1000行）"],
            ["generate_feature", "特征代码生成+注册", "名称+代码+维度+解释+知识溯源", "注册确认"],
            ["execute_sandbox", "三级验证", "特征名称", "L1/L2/L3结果+分数"],
            ["get_feature_status", "查询已有特征状态", "无", "有效特征列表+SHAP排名"],
            ["get_failure_memory", "查看失败记录", "无", "失败特征+原因+教训"],
        ],
        col_widths=[3.5, 3.5, 4, 4])

    doc.add_page_break()

    # ============================================================
    # 第5章 三级验证
    # ============================================================
    doc.add_heading("5. 验证特征有效性（三级验证）", level=1)

    add_image_or_fallback(doc, "sec5_verification.png", "三级验证详情", width=Inches(6.2))

    doc.add_heading("5.1 L1 沙盒执行", level=2)
    add_styled_paragraph(doc,
        "在隔离的子进程中执行特征代码，捕获异常和超时（15秒）。验证输出格式（DataFrame、列名、行数匹配）。",
        font_size=10.5)

    doc.add_heading("5.2 L2 统计质量（噪声校准）", level=2)
    add_styled_paragraph(doc,
        "核心创新：使用随机特征作为噪声基线，动态计算IV/KS阈值。\n"
        "生成5个随机特征 → 计算其IV/KS → 取95百分位作为阈值。\n"
        "只有超过随机噪声上界的特征才通过L2。\n"
        "同时检测冗余度（与已有特征的最大相关系数 < 0.7）和缺失率。",
        font_size=10.5)

    doc.add_heading("5.3 L3 模型效果", level=2)
    add_styled_paragraph(doc,
        "将新特征加入已有特征集，训练XGBoost模型，对比加入前后的delta指标：\n"
        "Delta Recall ≥ 0、Delta F1 ≥ 0、Precision ≥ 0.90。\n"
        "三者同时满足才通过L3。",
        font_size=10.5)

    doc.add_heading("5.4 结果分流与前向选择", level=2)
    add_styled_paragraph(doc,
        "L3通过 → 加入有效特征集\n"
        "L2通过但L3未通过 → 加入候选特征集，等待前向选择\n"
        "L2未通过 → 淘汰，记入FailureMemory\n\n"
        "每轮结束时对候选特征集执行贪心前向选择，恢复单独弱但组合强的特征。",
        font_size=10.5)

    doc.add_page_break()

    # ============================================================
    # 第6章 搜索策略
    # ============================================================
    doc.add_heading("6. Agent 的搜索策略", level=1)

    add_image_or_fallback(doc, "sec6_search_strategy.png", "搜索策略详情", width=Inches(6.2))

    doc.add_heading("6.1 维度化搜索", level=2)
    add_styled_paragraph(doc,
        "将特征空间按领域维度组织：金额频次、时间地域、行为序列、主体关联。\n"
        "每轮（Round）聚焦一个维度，预算为该维度的候选特征数（默认3-5个）。",
        font_size=10.5)

    doc.add_heading("6.2 SHAP反馈预算分配", level=2)
    add_styled_paragraph(doc,
        "每轮结束后计算SHAP特征重要性 → 按维度聚合 → 重要性低的维度获得更多预算。\n"
        "公式：budget_dim = base_budget × (1 + α × (1 - normalized_importance))\n"
        "确保探索覆盖所有维度，避免陷入局部最优。",
        font_size=10.5)

    doc.add_page_break()

    # ============================================================
    # 第7、8章
    # ============================================================
    doc.add_heading("7. Agent 的记忆与学习", level=1)
    add_image_or_fallback(doc, "sec7_memory.png", "Agent记忆与学习", width=Inches(6.2))

    doc.add_heading("7.1 FailureMemory", level=2)
    add_styled_paragraph(doc,
        "记录每个失败特征的：名称、失败阶段（L1/L2/L3）、失败原因、失败次数。\n"
        "注入System Prompt，防止Agent重复生成相同或相似的失败特征。",
        font_size=10.5)

    doc.add_heading("7.2 Provenance Card", level=2)
    add_styled_paragraph(doc,
        "每个通过验证的特征自动生成溯源卡片：\n"
        "• knowledge_ids: 引用的知识条目ID列表\n"
        "• reasoning: 设计推理（从知识到特征的逻辑链）\n"
        "• data_fields_used: 使用的数据字段\n"
        "• explanation: 自然语言解释",
        font_size=10.5)

    doc.add_page_break()

    doc.add_heading("8. 收敛检测与终止", level=1)
    add_image_or_fallback(doc, "sec8_convergence.png", "收敛检测", width=Inches(6.2))

    add_styled_paragraph(doc,
        "搜索策略循环图展示了Round N的完整流程：\n"
        "维度预算分配 → Agent生成候选 → 三级验证 → 前向选择 → SHAP反馈 → 更新预算 → 下一轮。\n"
        "最多10轮，当连续2轮无新增有效特征时提前终止。",
        font_size=10.5)

    doc.add_page_break()

    # ============================================================
    # 第9、10章
    # ============================================================
    doc.add_heading("9. 评价指标与评估方法", level=1)
    add_image_or_fallback(doc, "sec9_metrics.png", "评价指标", width=Inches(6.2))

    doc.add_heading("9.1 评估协议", level=2)
    make_table(doc,
        ["项目", "设置"],
        [
            ["划分方式", "80/20 train/test split"],
            ["重复次数", "5次不同随机种子 (42, 123, 456, 789, 2024)"],
            ["下游模型", "XGBoost (n_estimators=200, max_depth=6, lr=0.1)"],
            ["显著性检验", "Welch's t-test (p=0.05)"],
        ],
        col_widths=[4, 10])

    doc.add_heading("9.2 评估指标", level=2)
    make_table(doc,
        ["指标", "用途"],
        [
            ["Recall", "AML场景核心指标（漏检代价远高于误报）"],
            ["Precision", "约束条件（≥90%）"],
            ["F1", "综合指标"],
            ["AUC", "公开数据集主指标（与其他论文对齐）"],
        ],
        col_widths=[3, 11])

    doc.add_page_break()

    doc.add_heading("10. 与传统方法的对比", level=1)
    add_image_or_fallback(doc, "sec10_comparison.png", "与传统方法对比", width=Inches(6.2))

    make_table(doc,
        ["维度", "OpenFE", "CAAFE", "FeatLLM", "KnowFeat"],
        [
            ["方法类型", "数学组合搜索", "LLM生成", "LLM生成", "LLM+领域知识"],
            ["知识来源", "无", "数据集描述", "数据集描述", "结构化领域知识库"],
            ["可解释性", "无语义解释", "代码注释", "自然语言描述", "Provenance Card（知识溯源）"],
            ["知识可追溯", "✗", "✗", "✗", "✓ 精确到IND/RULE/EXP/RF"],
            ["验证机制", "统计筛选", "人工检查", "线性模型", "三级验证+噪声校准"],
            ["反馈机制", "无", "单轮", "单轮", "SHAP反馈+预算动态分配"],
            ["特征效率", "中（41.3个）", "中低（14.2个）", "低（148.2个）", "高（24.7个）"],
        ],
        col_widths=[3, 2.8, 2.8, 2.8, 3.5])

    doc.add_page_break()

    # ============================================================
    # 附录：交互示例
    # ============================================================
    doc.add_heading("附. 一次完整交互示例", level=1)
    add_styled_paragraph(doc,
        "以下展示 Round 1 金额频次维度中，Agent 生成两个候选特征的交互过程（一个通过、一个失败）。",
        font_size=10.5, space_after=8)
    add_image_or_fallback(doc, "sec_example.png", "完整交互示例", width=Inches(6.2))

    doc.add_page_break()

    # ============================================================
    # 实验结果：18个特征知识溯源卡片
    # ============================================================
    doc.add_heading("R. 实验结果：18个特征知识溯源卡片", level=1)

    add_styled_paragraph(doc,
        "以下是 KnowFeat 在数字人民币反洗钱(simecny)数据集上生成的18个有效特征。"
        "每个特征附带完整的知识溯源卡片（Provenance Card），精确引用领域知识资产编号。",
        font_size=10.5)

    add_styled_paragraph(doc,
        "增量效果：基线 Recall 0.9538 → 最终 0.9737 (+2.09%)，"
        "基线 Precision 0.9402 → 最终 0.9621 (+2.33%)。",
        font_size=10.5, bold=True, space_after=12)

    # 知识覆盖统计
    doc.add_heading("知识覆盖统计", level=2)
    make_table(doc,
        ["知识类型", "总数", "被引用数", "引用率"],
        [
            ["IND_xx (风险指标)", "38", "22", "57.9%"],
            ["RULE_xxx (检测规则)", "15", "10", "66.7%"],
            ["EXP_xx (专家意见)", "26", "10", "38.5%"],
            ["PAT_Axx (异常交易模式)", "5", "5", "100%"],
            ["RF_S2_xxx (裁判文书)", "~1920", "18", "~0.9%"],
        ],
        col_widths=[5, 2, 2.5, 2.5])

    doc.add_paragraph()

    # 加载 provenance 和 features
    with open(PROVENANCE_JSON, encoding="utf-8") as f:
        provenance_list = json.load(f)
    with open(FEATURES_JSON, encoding="utf-8") as f:
        features_list = json.load(f)

    # 按名字建索引
    code_map = {f["name"]: f["code"] for f in features_list}

    for idx, prov in enumerate(provenance_list, 1):
        fname = prov["feature_name"]
        dims = "、".join(prov["dimensions"])
        kids = prov.get("knowledge_ids", [])
        explanation = prov.get("explanation", "")
        reasoning = prov.get("reasoning", "")
        fields = prov.get("data_fields_used", [])
        selection = prov.get("selection_method", "")
        code = code_map.get(fname, "")

        # 标题
        sel_tag = "  [前向选择]" if selection else ""
        doc.add_heading(f"F{idx:02d}. {fname}{sel_tag}", level=2)

        # 维度和含义
        p = doc.add_paragraph()
        run = p.add_run(f"维度: {dims}")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(13, 148, 136)

        if explanation:
            add_styled_paragraph(doc, f"含义: {explanation}", font_size=10)

        # 知识溯源表
        if kids:
            p = doc.add_paragraph()
            run = p.add_run("知识溯源:")
            run.bold = True
            run.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(2)

            # 简单列出ID
            id_text = "  ".join(kids)
            p2 = doc.add_paragraph()
            run2 = p2.add_run(id_text)
            run2.font.size = Pt(9)
            run2.font.name = "Consolas"
            run2.font.color.rgb = RGBColor(26, 86, 219)
            p2.paragraph_format.space_after = Pt(4)

        # 设计推理
        if reasoning:
            p = doc.add_paragraph()
            run_label = p.add_run("设计推理: ")
            run_label.bold = True
            run_label.font.size = Pt(9.5)
            run_text = p.add_run(reasoning)
            run_text.font.size = Pt(9.5)
            p.paragraph_format.space_after = Pt(4)

        # 数据字段
        if fields:
            p = doc.add_paragraph()
            run = p.add_run(f"数据字段: {', '.join(fields)}")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 116, 139)

        # Python代码
        if code:
            p = doc.add_paragraph()
            run = p.add_run("Python 代码:")
            run.bold = True
            run.font.size = Pt(9.5)
            p.paragraph_format.space_after = Pt(2)
            add_code_block(doc, code.strip())

        # 分隔线（空段落 + 细线）
        if idx < len(provenance_list):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="E2E8F0"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)

    # ============================================================
    # 页脚
    # ============================================================
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("KnowFeat: Knowledge-Traceable Feature Engineering — 技术方案文档 v4")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 116, 139)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = doc.add_paragraph()
    run2 = p2.add_run("Generated: 2026-06-01")
    run2.font.size = Pt(8)
    run2.font.color.rgb = RGBColor(100, 116, 139)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 保存
    doc.save(str(OUTPUT_PATH))
    print(f"Word document saved to: {OUTPUT_PATH}")
    print(f"File size: {OUTPUT_PATH.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    build_document()
